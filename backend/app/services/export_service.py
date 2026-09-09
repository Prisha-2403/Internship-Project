"""CSV and Excel export of transaction data.

Exports honour the same masking as the API: a spreadsheet leaving the platform
must not contain anything the screen would have hidden.

CSV output is guarded against formula injection - a cell beginning with ``=``,
``+``, ``-`` or ``@`` is prefixed so a spreadsheet treats it as text rather
than executing it on open.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.masking import mask_device_ref, mask_name

COLUMNS: tuple[tuple[str, str], ...] = (
    ("transaction_ref", "Transaction ID"),
    ("customer_ref", "Customer ID"),
    ("customer_name", "Customer"),
    ("amount", "Amount"),
    ("currency", "Currency"),
    ("occurred_at", "Date/Time"),
    ("location_city", "Location"),
    ("device_ref", "Device"),
    ("payment_method", "Payment Method"),
    ("risk_score", "Risk Score"),
    ("risk_level", "Risk Level"),
    ("primary_reason", "Detection Reason"),
    ("status", "Status"),
)

#: Characters a spreadsheet may interpret as the start of a formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitise(value: Any) -> Any:
    """Neutralise spreadsheet formula injection in text cells."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def build_rows(records: list[tuple]) -> list[dict[str, Any]]:
    """Flatten query rows into export records, applying masking."""
    rows: list[dict[str, Any]] = []
    for transaction, risk, customer, device in records:
        rows.append(
            {
                "transaction_ref": transaction.transaction_ref,
                "customer_ref": customer.customer_ref,
                "customer_name": mask_name(customer.full_name),
                "amount": float(transaction.amount),
                "currency": transaction.currency,
                "occurred_at": transaction.occurred_at.strftime("%Y-%m-%d %H:%M:%S"),
                "location_city": transaction.location_city,
                "device_ref": mask_device_ref(device.device_ref) if device else "",
                "payment_method": transaction.payment_method.value,
                "risk_score": risk.business_score if risk else "",
                "risk_level": risk.risk_level.value if risk else "",
                "primary_reason": (risk.primary_reason if risk else "") or "",
                "status": transaction.status.value,
            }
        )
    return rows


def to_csv(rows: list[dict[str, Any]]) -> bytes:
    """Render rows as UTF-8 CSV with a BOM so Excel opens it correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow([header for _, header in COLUMNS])
    for row in rows:
        writer.writerow([_sanitise(row.get(key, "")) for key, _ in COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict[str, Any]], title: str = "Transactions") -> bytes:
    """Render rows as a formatted single-sheet workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E293B")
    for column_index, (_, header) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _) in enumerate(COLUMNS, start=1):
            value = row.get(key, "")
            cell = sheet.cell(row=row_index, column=column_index, value=_sanitise(value))
            if key == "amount":
                cell.number_format = "#,##0.00"

    widths = {
        "transaction_ref": 16, "customer_ref": 14, "customer_name": 20, "amount": 14,
        "currency": 10, "occurred_at": 20, "location_city": 16, "device_ref": 18,
        "payment_method": 16, "risk_score": 12, "risk_level": 12,
        "primary_reason": 22, "status": 14,
    }
    for column_index, (key, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(key, 16)

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(extension: str, prefix: str = "sentinel-transactions") -> str:
    return f"{prefix}-{datetime.now(UTC):%Y%m%d-%H%M%S}.{extension}"


MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
